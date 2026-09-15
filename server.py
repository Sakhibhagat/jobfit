import os
import re
from io import BytesIO
from dotenv import load_dotenv
from groq import Groq
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt, RGBColor
import fitz  # PyMuPDF
import markdown
from flask import Flask, render_template, request, send_file, jsonify

load_dotenv()
client = Groq(api_key=os.getenv("GROQ_API_KEY"))
app = Flask(__name__)

session_cache = {}


def call_ai(prompt):
    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def extract_text_from_pdf(file_stream):
    reader = PdfReader(file_stream)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text


def extract_style_info(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    font_counts = {}
    heading_color = None
    max_size_seen = 0

    for page in doc:
        blocks = page.get_text("dict")["blocks"]
        for b in blocks:
            for l in b.get("lines", []):
                for s in l.get("spans", []):
                    font = s.get("font", "")
                    size = s.get("size", 0)
                    color = s.get("color", 0)
                    text_len = len(s.get("text", "").strip())
                    if text_len == 0:
                        continue
                    font_counts[font] = font_counts.get(font, 0) + text_len
                    if size > max_size_seen and color != 0:
                        max_size_seen = size
                        heading_color = color

    doc.close()

    dominant_font = max(font_counts, key=font_counts.get) if font_counts else "Calibri"
    clean_font = dominant_font.split("+")[-1] if "+" in dominant_font else dominant_font
    for suffix in ["-Bold", "-Italic", "-Regular", "-BoldItalic", ",Bold", ",Italic"]:
        clean_font = clean_font.replace(suffix, "")

    def int_to_hex(c):
        if not c:
            return "#1C1B19"
        r = (c >> 16) & 255
        g = (c >> 8) & 255
        b = c & 255
        return f"#{r:02x}{g:02x}{b:02x}"

    return {"font": clean_font.strip() or "Calibri", "color": int_to_hex(heading_color)}


def get_feedback(resume_text, job_description):
    prompt = f"""You are an expert resume reviewer. Analyze this resume against the job description below.
Provide:
1. Overall fit assessment (2-3 sentences)
2. Top 3 strengths relevant to this job
3. Top 3 gaps or areas to improve for this specific role

Job Description:
{job_description}

Resume:
{resume_text}
"""
    return call_ai(prompt)


def get_updated_resume(resume_text, job_description, additional_info=""):
    prompt = f"""You are an expert resume writer. Rewrite the resume below to be more targeted for the job description provided.
Fix grammar, weak phrasing, and vague bullet points. Emphasize relevant skills and experience for this specific job.
Keep the person's real experience and facts — do not invent new jobs, companies, or degrees.

Format your output EXACTLY like this structure, using these exact markers:
##NAME##
(person's name)
##SECTION##Contact
(contact info)
##SECTION##Summary
(2-3 sentence professional summary tailored to this job)
##SECTION##Experience
(job entries, each bullet starting with "- ")
##SECTION##Education
(education entries)
##SECTION##Skills
(skills, comma separated or bullets)

Job Description:
{job_description}

Additional instructions from the candidate (apply these carefully):
{additional_info if additional_info.strip() else "None provided."}

Original Resume:
{resume_text}
"""
    return call_ai(prompt)


def revise_resume(previous_draft, revision_notes):
    prompt = f"""Here is a resume draft using this format (keep the exact same format in your response):
##NAME##
Name
##SECTION##Header
content
(and so on for each section)

Current draft:
{previous_draft}

The candidate wants these changes made:
{revision_notes}

Return the FULL revised resume using the exact same ##NAME##/##SECTION## marker format, with the requested changes applied."""
    return call_ai(prompt)


def structured_to_html(tailored_text, style_info=None):
    font = (style_info or {}).get("font", "Georgia")
    color = (style_info or {}).get("color", "#1C1B19")

    name_match = re.search(r"##NAME##\s*(.*?)\s*##SECTION##", tailored_text, re.DOTALL)
    name = name_match.group(1).strip() if name_match else ""

    html = f'<div class="resume-doc" style="font-family: \'{font}\', Georgia, serif;">'
    if name:
        html += f'<h1 style="color: {color}; border-color: {color};">{name}</h1>'

    sections = re.split(r"##SECTION##", tailored_text)[1:]
    for section in sections:
        lines = section.strip().split("\n", 1)
        heading = lines[0].strip()
        body = lines[1].strip() if len(lines) > 1 else ""

        html += f'<h2 style="color: {color};">{heading}</h2>'
        for line in body.split("\n"):
            line = line.strip()
            if not line:
                continue
            clean_line = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", line)
            clean_line = re.sub(r"\*(.*?)\*", r"<em>\1</em>", clean_line)
            if clean_line.startswith("- "):
                html += f'<p class="bullet">{clean_line[2:]}</p>'
            else:
                html += f'<p>{clean_line}</p>'

    html += '</div>'
    return html


def add_formatted_paragraph(doc, text, font_name, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            run = p.add_run(part)
        run.font.name = font_name
    return p


def build_docx(tailored_text, style_info=None):
    font_name = (style_info or {}).get("font", "Calibri")
    color_hex = (style_info or {}).get("color", "#1C1B19").lstrip("#")
    rgb = RGBColor(int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16))

    doc = Document()
    doc.styles["Normal"].font.name = font_name
    doc.styles["Normal"].font.size = Pt(11)

    name_match = re.search(r"##NAME##\s*(.*?)\s*##SECTION##", tailored_text, re.DOTALL)
    name = name_match.group(1).strip() if name_match else "Resume"
    heading = doc.add_heading(name, level=0)
    for run in heading.runs:
        run.font.color.rgb = rgb
        run.font.name = font_name

    sections = re.split(r"##SECTION##", tailored_text)[1:]
    for section in sections:
        lines = section.strip().split("\n", 1)
        heading_text = lines[0].strip()
        body = lines[1].strip() if len(lines) > 1 else ""

        h = doc.add_heading(heading_text, level=1)
        for run in h.runs:
            run.font.color.rgb = rgb
            run.font.name = font_name

        for line in body.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("- "):
                add_formatted_paragraph(doc, line[2:], font_name, style="List Bullet")
            else:
                add_formatted_paragraph(doc, line, font_name)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/generate", methods=["POST"])
def generate():
    uploaded_file = request.files.get("resume")
    job_description = request.form.get("job_description", "")
    additional_info = request.form.get("additional_info", "")

    if not uploaded_file or not job_description.strip():
        return jsonify({"error": "Missing resume or job description"}), 400

    pdf_bytes = uploaded_file.read()
    resume_text = extract_text_from_pdf(BytesIO(pdf_bytes))
    style_info = extract_style_info(pdf_bytes)

    feedback_raw = get_feedback(resume_text, job_description)
    updated = get_updated_resume(resume_text, job_description, additional_info)

    session_cache["updated_resume"] = updated
    session_cache["style_info"] = style_info

    feedback_html = markdown.markdown(feedback_raw, extensions=["tables"])

    return jsonify({
        "feedback": feedback_html,
        "preview": structured_to_html(updated, style_info)
    })


@app.route("/revise", methods=["POST"])
def revise():
    revision_notes = request.form.get("revision_notes", "")
    previous = session_cache.get("updated_resume")
    style_info = session_cache.get("style_info")

    if not previous:
        return jsonify({"error": "No resume to revise yet"}), 400
    if not revision_notes.strip():
        return jsonify({"error": "Tell us what to change"}), 400

    updated = revise_resume(previous, revision_notes)
    session_cache["updated_resume"] = updated

    return jsonify({"preview": structured_to_html(updated, style_info)})


@app.route("/download")
def download():
    updated = session_cache.get("updated_resume")
    style_info = session_cache.get("style_info")
    if not updated:
        return "No resume generated yet", 400
    docx_buffer = build_docx(updated, style_info)
    return send_file(
        docx_buffer,
        as_attachment=True,
        download_name="updated_resume.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


if __name__ == "__main__":
    app.run(debug=True)