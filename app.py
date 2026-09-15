import os
import re
from io import BytesIO
from dotenv import load_dotenv
from groq import Groq
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt
import streamlit as st

load_dotenv()
client = Groq(api_key=os.getenv("GROQ_API_KEY"))

st.set_page_config(page_title="Resume Analyzer", page_icon="📄", layout="centered")

st.markdown("""
    <style>
    .stButton>button {
        background-color: #4F46E5;
        color: white;
        border-radius: 8px;
        padding: 0.5rem 1.5rem;
        font-weight: 600;
        border: none;
    }
    .stButton>button:hover { background-color: #4338CA; }
    .title-text {
        font-size: 2.5rem;
        font-weight: 800;
        background: linear-gradient(90deg, #4F46E5, #818CF8);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0;
    }
    .subtitle-text { color: #9CA3AF; font-size: 1.1rem; margin-top: 0; margin-bottom: 2rem; }
    .feedback-box {
        background-color: #1F2937;
        border-radius: 12px;
        padding: 1.5rem;
        border-left: 4px solid #4F46E5;
    }
    </style>
""", unsafe_allow_html=True)


def extract_text_from_pdf(uploaded_file):
    reader = PdfReader(uploaded_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text


def get_feedback(resume_text, job_description):
    prompt = f"""You are an expert resume reviewer. Analyze this resume against the job description below.
Provide:
1. Overall fit assessment (2-3 sentences)
2. Top 3 strengths relevant to this job
3. Top 3 gaps or areas to improve for this specific role

Job Description:
{job_description}

Resume:
{resume_text}
"""
    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def get_tailored_resume(resume_text, job_description):
    prompt = f"""You are an expert resume writer. Rewrite the resume below to be more targeted for the job description provided.
Fix grammar, weak phrasing, and vague bullet points. Emphasize relevant skills and experience for this specific job.
Keep the person's real experience and facts — do not invent new jobs, companies, or degrees.

Format your output EXACTLY like this structure, using these exact markers:
##NAME##
(person's name)
##SECTION##Contact
(contact info)
##SECTION##Summary
(2-3 sentence professional summary tailored to this job)
##SECTION##Experience
(job entries, each bullet starting with "- ")
##SECTION##Education
(education entries)
##SECTION##Skills
(skills, comma separated or bullets)

Job Description:
{job_description}

Original Resume:
{resume_text}
"""
    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content


def add_formatted_paragraph(doc, text, style=None):
    """Adds a paragraph, converting **bold** and *italic* markdown into real formatting."""
    if style:
        p = doc.add_paragraph(style=style)
    else:
        p = doc.add_paragraph()

    # Split on **bold** and *italic* markers, keeping the markers to know which parts to format
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            p.add_run(part)
    return p


def build_docx(tailored_text):
    doc = Document()

    name_match = re.search(r"##NAME##\s*(.*?)\s*##SECTION##", tailored_text, re.DOTALL)
    name = name_match.group(1).strip() if name_match else "Resume"

    doc.add_heading(name, level=0)

    sections = re.split(r"##SECTION##", tailored_text)[1:]
    for section in sections:
        lines = section.strip().split("\n", 1)
        heading = lines[0].strip()
        body = lines[1].strip() if len(lines) > 1 else ""

        doc.add_heading(heading, level=1)
        for line in body.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("- "):
                add_formatted_paragraph(doc, line[2:], style="List Bullet")
            else:
                add_formatted_paragraph(doc, line)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer




st.markdown('<p class="title-text">📄 Resume Analyzer</p>', unsafe_allow_html=True)
st.markdown('<p class="subtitle-text">Get AI feedback and a tailored resume for the job you want</p>', unsafe_allow_html=True)

uploaded_file = st.file_uploader("Upload your resume (PDF)", type=["pdf"])
job_description = st.text_area("Paste the job description you're applying for", height=200)

if uploaded_file is not None and job_description.strip():
    if st.button("Analyze & Tailor My Resume →"):
        with st.spinner("Reading your resume..."):
            resume_text = extract_text_from_pdf(uploaded_file)

        with st.spinner("Analyzing fit..."):
            feedback = get_feedback(resume_text, job_description)

        with st.spinner("Rewriting your resume for this job..."):
            tailored = get_tailored_resume(resume_text, job_description)

        st.markdown("---")
        st.markdown("### ✨ Feedback")
        st.markdown(f'<div class="feedback-box">{feedback}</div>', unsafe_allow_html=True)

        st.markdown("### 📝 Tailored Resume")
        docx_buffer = build_docx(tailored)
        st.download_button(
            label="Download Tailored Resume (.docx)",
            data=docx_buffer,
            file_name="tailored_resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
else:
    st.info("👆 Upload your resume and paste a job description to get started")